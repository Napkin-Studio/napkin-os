"""Offline tests for the native-shape layout library and the theme (no model, no network)."""
from __future__ import annotations

import copy
import json
import re
import shutil
import subprocess
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Emu

from deck import render_pptx as R
from deck import theme as T
from deck.inputs import Inputs
from deck.spec import shown_blocks

FIX = Path(__file__).parent / "fixtures" / "lumen_water"
SOFFICE = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")


def _deck():
    """The golden brief-only deck (invented content)."""
    return json.loads((FIX / "golden_deck.json").read_text())


def _inp():
    """Empty inputs: the golden deck has no chart."""
    return Inputs(brief={})


def _render(tmp_path, deck=None, **kw):
    """Render a deck dict; returns (Renderer, Presentation)."""
    deck = deck or _deck()
    out = tmp_path / "d.pptx"
    r = R.Renderer(deck, _inp(), theme=kw.pop("theme", "napkin"), **kw)
    r.render(out)
    return r, Presentation(str(out))


def norm(t):
    """Whitespace-normalised text."""
    return re.sub(r"\s+", " ", str(t)).strip()


def slide_text(sl):
    """All text on a slide (shapes in order, paragraphs and table cells joined by spaces), normalised."""
    out = []
    for sh in sl.shapes:
        if sh.has_text_frame:
            out.append(" ".join(p.text for p in sh.text_frame.paragraphs))
        elif getattr(sh, "has_table", False) and sh.has_table:
            for row in sh.table.rows:
                out.extend(c.text for c in row.cells)
    return norm(" ".join(out))


def strings(b):
    """Every string a block carries on the slide."""
    k = b["kind"]
    if k in ("text", "small", "statement", "quote", "callout"):
        return [b["text"]]
    if k == "card":
        return [b.get("title", ""), b["text"]]
    if k == "bullets":
        return ([b["title"]] if b.get("title") else []) + [i["text"] if isinstance(i, dict) else i for i in b["items"]]
    if k == "table":
        return [c for r in b["rows"] for c in r]
    if k == "kv":
        return [y for k, v in b["items"] for y in [k] + (list(v) if isinstance(v, (list, tuple)) else [v])]
    if k == "fromto":
        return [b["from"], b["to"]]
    if k == "stats":
        return [x for i in b["items"] for x in (i["value"], i["label"])]
    if k == "steps":
        return [x for i in b["items"] for x in i]
    if k == "copy":
        return list(b["lines"]) + [b["answer"]]
    if k == "sketch":
        return [t for t in re.split(r"\s*->\s*", b["text"]) if t]
    return []


def test_every_word_survives(tmp_path):
    deck = _deck()
    r, prs = _render(tmp_path, deck)
    for sid, sl in zip(deck["order"], prs.slides):
        s = deck["slides"][sid]
        text = slide_text(sl)
        want = [s["title"]] + [x for b in shown_blocks(s) for x in strings(b)]
        for w in want:
            if w.strip():
                assert norm(w) in text, (sid, r.layouts[sid], w)


def test_claims_are_kept_in_the_file(tmp_path):
    """Claims are evidence: each is on its slide or in its notes."""
    deck = _deck()
    r, prs = _render(tmp_path, deck)
    for sid, sl in zip(deck["order"], prs.slides):
        have = slide_text(sl) + " " + norm(sl.notes_slide.notes_text_frame.text)
        for c in deck["slides"][sid]["claims"]:
            assert norm(c["text"]) in have, (sid, c["text"])


def test_nothing_past_the_slide_and_no_small_type(tmp_path):
    r, prs = _render(tmp_path)
    W, H = prs.slide_width, prs.slide_height
    for i, sl in enumerate(prs.slides, 1):
        for sh in sl.shapes:
            assert sh.left >= 0 and sh.top >= 0, (i, sh.name)
            assert sh.left + sh.width <= W + 10, (i, sh.name, "right")
            assert sh.top + sh.height <= H + 10, (i, sh.name, "bottom")
            frames = [sh.text_frame] if sh.has_text_frame else []
            if getattr(sh, "has_table", False) and sh.has_table:
                frames = [c.text_frame for row in sh.table.rows for c in row.cells]
            for tf in frames:
                for p in tf.paragraphs:
                    for run in p.runs:
                        if run.text.strip():
                            assert run.font.size is not None and run.font.size.pt >= 10, (i, sh.name, run.font.size)


def test_accent_ink_passes_on_surface_and_panel():
    t = T.get_theme("napkin")
    assert T.contrast(t.accent_ink, t.paper) >= 4.5
    assert T.contrast(t.accent_ink, t.panel) >= 4.5
    assert T.check_palette(t) == []
    assert 36 <= t.scale.title <= 44 and 14 <= t.scale.body <= 16 and 10 <= t.scale.caption <= 12
    assert t.grid.M >= 0.5


def test_contrast_function_known_values():
    assert round(T.contrast("000000", "FFFFFF"), 1) == 21.0
    assert T.contrast("777777", "777777") == pytest.approx(1.0)


def test_check_palette_reports_problems():
    bad = T.get_theme("napkin", {"colors": {"accent_ink": "F0EADD"}})
    assert any("accent_ink on paper" in p for p in T.check_palette(bad))
    assert T.check_palette(T.CLASSIC)   # the old amber text is below 4.5:1, said plainly


def test_layout_chosen_per_type_as_documented(tmp_path):
    deck = _deck()
    r, _ = _render(tmp_path, deck)
    want = {"s01": "cover", "s03": "emphasis_bar", "s04": "quote_cards", "s06": "comparison", "s08": "emphasis_bar",
            "s09": "execution", "s10": "role_grid", "s12": "native_table", "s13": "timeline_strip",
            "s15": "placeholder_card", "s16": "placeholder_card", "s19": "territory", "s11": "cards"}
    for sid, name in want.items():
        assert R.choose_layout(deck["slides"][sid]) == name, sid
        assert r.layouts[sid] == name, (sid, "fell back")


def test_choose_layout_rules_on_synthetic_slides():
    def sl(t, *blocks, status="drafted"):
        return {"type": t, "status": status, "title": "T", "blocks": list(blocks), "claims": []}
    stat = lambda n: {"kind": "stats", "items": [{"value": "1", "label": "x"}] * n}
    assert R.choose_layout(sl("market", stat(1))) == "big_number"
    assert R.choose_layout(sl("market", stat(3))) == "kpi_row"
    assert R.choose_layout(sl("team", {"kind": "bullets", "items": [{"text": "a", "human": True}]})) == "placeholder_card"
    assert R.choose_layout(sl("team", {"kind": "text", "text": "x"}, status="needs_human")) == "placeholder_card"
    assert R.choose_layout(sl("founders", {"kind": "steps", "items": [["a", "b"]]})) == "role_grid"
    assert R.choose_layout(sl("weird", {"kind": "mystery"})) == "flow"


def test_fallback_on_an_odd_slide(tmp_path):
    deck = _deck()
    deck["slides"]["s05"]["blocks"].append({"kind": "mystery", "text": "odd"})
    r, prs = _render(tmp_path, deck)
    assert r.layouts["s05"] == "flow"
    assert len(prs.slides) == len(deck["order"])


def test_fallback_on_a_slide_too_dense(tmp_path):
    deck = _deck()
    deck["slides"]["s11"]["blocks"] = [{"kind": "bullets", "items": [{"text": "word " * 120} for _ in range(14)]}]
    r, prs = _render(tmp_path, deck)
    assert r.layouts["s11"] == "flow"
    assert "word" in slide_text(prs.slides[10])


def test_classic_theme_still_renders_with_the_flow(tmp_path):
    r, prs = _render(tmp_path, theme="classic")
    assert set(r.layouts.values()) == {"flow"}
    assert r.th.paper == "FFFFFF" and r.th.accent == "E8A33D"
    assert len(prs.slides) == len(_deck()["order"])


def test_unknown_theme_is_an_error(tmp_path):
    with pytest.raises(ValueError):
        R.Renderer(_deck(), _inp(), theme="nope")


def test_brand_json_still_applies(tmp_path):
    deck = _deck()
    deck["meta"]["brand"] = {"colors": {"ink": "#102030", "accent": "#00AA55", "paper": "#FFFFFF"}, "head": "Georgia", "body": "Arial"}
    r, prs = _render(tmp_path, deck)
    assert r.th.ink == "102030" and r.th.accent == "00AA55" and r.th.head == "Georgia" and r.th.body == "Arial"
    assert T.contrast(r.th.accent_ink, r.th.paper) >= 4.5   # derived, so text stays legible
    xml = prs.slides[1]._element.xml
    assert "102030" in xml                                    # the dark slide's background
    fonts = {run.font.name for sl in prs.slides for sh in sl.shapes if sh.has_text_frame for p in sh.text_frame.paragraphs for run in p.runs}
    assert fonts <= {"Georgia", "Arial"}


def test_fonts_are_the_ones_already_named(tmp_path):
    r, prs = _render(tmp_path)
    fonts = {run.font.name for sl in prs.slides for sh in sl.shapes if sh.has_text_frame for p in sh.text_frame.paragraphs for run in p.runs}
    assert fonts <= {"Cambria", "Calibri"}


def test_notes_sections_watermark_and_hidden_appendix_kept(tmp_path):
    deck = _deck()
    r, prs = _render(tmp_path, deck)
    for sid, sl in zip(deck["order"], prs.slides):
        assert f"Slide id {sid}" in sl.notes_slide.notes_text_frame.text
        assert "DRAFT" in slide_text(sl)
    shown = [sl._element.get("show") for sl in prs.slides]
    first = deck["order"].index("s19")
    assert shown[first] == "0" and shown[-1] == "0" and shown[0] is None
    assert "sectionLst" in prs.part._element.xml


def test_animation_option_still_builds(tmp_path):
    r, prs = _render(tmp_path, animate="auto")
    assert "<p:timing" in prs.slides[2]._element.xml and "<p:transition" in prs.slides[2]._element.xml


def test_placeholder_is_never_styled_as_finished(tmp_path):
    deck = _deck()
    r, prs = _render(tmp_path, deck)
    sl = prs.slides[deck["order"].index("s15")]
    text = slide_text(sl)
    assert R.CHIP in text and R.TAG in text
    assert "prstDash" in sl._element.xml


def test_render_module_entry_point_and_render_web_import(tmp_path):
    out = R.render(_deck(), _inp(), tmp_path / "x.pptx")
    assert out.exists()
    from deck import render_web
    assert render_web.THEME is R.THEME


@pytest.mark.skipif(not SOFFICE.exists(), reason="LibreOffice not installed")
def test_libreoffice_converts_a_render():
    import tempfile
    with tempfile.TemporaryDirectory(dir="/tmp") as d:
        out = Path(d) / "d.pptx"
        R.render(_deck(), _inp(), out)
        subprocess.run([str(SOFFICE), "--headless", "--convert-to", "pdf", "--outdir", d, str(out)], check=True, capture_output=True, timeout=180)
        assert (Path(d) / "d.pdf").stat().st_size > 1000


def test_kv_list_values_render_as_words_not_a_python_list(tmp_path):
    deck = _deck()
    deck["slides"]["s19"]["blocks"][1] = {"kind": "kv", "items": [["Risks", ["First risk here.", "Second risk here."]]]}
    r, prs = _render(tmp_path, deck)
    text = slide_text(prs.slides[deck["order"].index("s19")])
    assert "First risk here." in text and "Second risk here." in text and "['" not in text


def _human_ids(deck):
    """Slide ids whose bullets carry a human-flagged item."""
    return [sid for sid, s in deck["slides"].items()
            if any(isinstance(i, dict) and i.get("human") for b in shown_blocks(s) if b["kind"] == "bullets" for i in b["items"])]


def test_placeholder_marks_survive_the_classic_theme(tmp_path):
    """Human bullets keep the dashed card, PROPOSED tag and TO SUPPLY chip under the flow-only classic theme."""
    deck = _deck()
    ids = _human_ids(deck)
    assert ids
    r, prs = _render(tmp_path, deck, theme="classic")
    for sid in ids:
        sl = prs.slides[deck["order"].index(sid)]
        assert r.layouts[sid] == "flow"
        assert R.TAG in slide_text(sl) and R.CHIP in slide_text(sl), sid
        assert "prstDash" in sl._element.xml, sid
        for b in shown_blocks(deck["slides"][sid]):
            for w in strings(b):
                assert norm(w) in slide_text(sl), (sid, w)


def test_placeholder_marks_survive_a_layout_fallback(tmp_path, monkeypatch):
    """When placeholder_card raises NoFit the earlier flow still marks every human bullet."""
    def nofit(self, sl, s):
        """Stand-in layout that never fits."""
        raise R.NoFit("forced")
    monkeypatch.setitem(R.Renderer.LAYOUT_FN, "placeholder_card", "_nofit_layout")
    monkeypatch.setattr(R.Renderer, "_nofit_layout", nofit, raising=False)
    deck = _deck()
    r, prs = _render(tmp_path, deck)
    for sid in _human_ids(deck):
        assert r.layouts[sid] == "flow"
        sl = prs.slides[deck["order"].index(sid)]
        n_human = sum(1 for b in shown_blocks(deck["slides"][sid]) if b["kind"] == "bullets"
                      for i in b["items"] if isinstance(i, dict) and i.get("human"))
        assert slide_text(sl).count(R.CHIP) == n_human and slide_text(sl).count(R.TAG) == n_human, sid


def _png() -> bytes:
    """A valid 2x1 RGB PNG built with the stdlib (no image library, no binary fixture)."""
    import struct
    import zlib

    def chunk(tag, data):
        """One PNG chunk with its CRC."""
        return struct.pack(">I", len(data)) + tag + data + struct.pack(">I", zlib.crc32(tag + data) & 0xFFFFFFFF)
    raw = b"\x00" + b"\xff\x00\x00" + b"\x00\x00\xff"
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", 2, 1, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))


def test_logo_is_the_brands_own_with_alt_text_and_the_kicker_stays(tmp_path):
    """A brand logo on the cover carries alt text, and the kicker words are kept beside it (both paths)."""
    logo = tmp_path / "logo.png"
    logo.write_bytes(_png())
    deck = _deck()
    deck["meta"]["brand"] = {"logo_dark_png": str(logo)}
    r, prs = _render(tmp_path, deck)
    sl = prs.slides[0]
    pics = [sh for sh in sl.shapes if sh.shape_type == 13]
    assert len(pics) == 1 and pics[0]._element.nvPicPr.cNvPr.get("descr")
    assert norm(R._kicker(deck["meta"]).upper()) in slide_text(sl)
    r2, prs2 = _render(tmp_path, deck, theme="classic")
    sl2 = prs2.slides[0]
    assert any(sh.shape_type == 13 and sh._element.nvPicPr.cNvPr.get("descr") for sh in sl2.shapes)
    assert norm(R._kicker(deck["meta"]).upper()) in slide_text(sl2)


# ---------------------------------------------------------------------------------------------
# Balance round: the height is used on purpose, type and tables are even, floors agree (all offline)
# ---------------------------------------------------------------------------------------------
from deck import qa_render  # noqa: E402

IN = 914400


def _slide_deck(slide_type, blocks, title="A short claim about the water", section="Insight"):
    """The golden deck cut down to one slide of the given type and blocks (invented words)."""
    deck = _deck()
    deck["slides"] = {"sx": {"type": slide_type, "section": section, "title": title, "label": title, "blocks": blocks,
                             "claims": [], "visual": {"kind": "layout"}, "notes": "", "status": "drafted",
                             "requirements": [], "footnotes": []}}
    deck["order"] = ["sx"]
    return deck


def _fills(sl, hex_):
    """Plain filled rectangles of one colour on a slide, as (left, top, width, height) in inches, by top then left."""
    out = []
    for sh in sl.shapes:
        if sh.shape_type == 1 and sh.fill.type == 1 and str(sh.fill.fore_color.rgb) == hex_:
            out.append((sh.left / IN, sh.top / IN, sh.width / IN, sh.height / IN))
    return sorted(out, key=lambda r: (round(r[1], 2), r[0]))


def _cards(sl, th):
    """The card rectangles (panel fill, tall enough to be a card, not an accent strip)."""
    return [r for r in _fills(sl, th.panel) if r[3] > 0.6]


SHORT = [{"kind": "bullets", "items": [{"text": "A short claim."}, {"text": "A second short claim."}, {"text": "A third one."}]}]


def test_short_content_cards_reach_the_foot_of_the_content_area(tmp_path):
    """A slide with little content stretches its cards to the bottom of the content area (6.75 in)."""
    r, prs = _render(tmp_path, _slide_deck("insight", SHORT))
    assert r.layouts["sx"] == "cards"
    cards = _cards(prs.slides[0], r.th)
    assert len(cards) == 3
    bottoms = [c[1] + c[3] for c in cards]
    assert all(abs(b - r.th.grid.body_end) < 0.08 for b in bottoms), bottoms


def test_paired_cards_in_a_row_share_one_height(tmp_path):
    """Cards of a row are one height even when their text is not."""
    blocks = [{"kind": "bullets", "items": [{"text": "Short."}, {"text": "A much longer claim that runs over several lines of the card so it is taller than its neighbours."},
                                           {"text": "Medium length claim here."}]}]
    r, prs = _render(tmp_path, _slide_deck("insight", blocks))
    cards = _cards(prs.slides[0], r.th)
    rows = {}
    for c in cards:
        rows.setdefault(round(c[1], 2), []).append(round(c[3], 2))
    assert rows and all(len(set(h)) == 1 for h in rows.values()), rows


def test_stretch_applies_to_every_golden_card_slide(tmp_path):
    """On the golden deck each plain card row ends at the content area's foot or where its stack ends: never below it."""
    r, prs = _render(tmp_path)
    for sl in prs.slides:
        for c in _cards(sl, r.th):
            assert c[1] + c[3] <= r.th.grid.body_end + 0.05


def test_dark_panel_and_its_neighbours_end_on_one_line(tmp_path):
    """Quote slide: the dark quote panel and the card beside it have the same top and the same bottom."""
    deck = _slide_deck("insight", [{"kind": "quote", "text": "A short quote."}, {"kind": "small", "text": "Illustrative."},
                                   {"kind": "bullets", "items": [{"text": "One bullet only."}]}])
    r, prs = _render(tmp_path, deck)
    assert r.layouts["sx"] == "quote_cards"
    sl = prs.slides[0]
    dark = [x for x in _fills(sl, r.th.ink) if x[2] > 3 and x[3] > 2]
    bar = [x for x in _fills(sl, r.th.accent) if x[3] > 2]
    assert dark and bar
    assert abs(dark[0][1] - bar[0][1]) < 0.02 and abs((dark[0][1] + dark[0][3]) - (bar[0][1] + bar[0][3])) < 0.05


def test_execution_closing_line_sits_a_fixed_gap_under_the_copy(tmp_path):
    """The closing answer is 0.4 in under the last line of the copy, inside the dark frame, and the light notes card
    ends on the same line as the frame."""
    cp = {"kind": "copy", "lines": ["One line.", "Another line."], "answer": "Because it answers the question."}
    deck = _slide_deck("execution", [{"kind": "text", "text": "Role: awareness."}, cp], section="The work")
    r, prs = _render(tmp_path, deck)
    assert r.layouts["sx"] == "execution"
    sl = prs.slides[0]
    frame = [x for x in _fills(sl, r.th.ink) if x[2] > 5]
    card = [x for x in _fills(sl, r.th.panel) if x[2] > 4 and x[3] > 2]
    assert frame and card
    assert abs((frame[0][1] + frame[0][3]) - (card[0][1] + card[0][3])) < 0.02
    boxes = {sh.text_frame.text.strip(): sh for sh in sl.shapes if sh.has_text_frame}
    copy_box = next(sh for t, sh in boxes.items() if t.startswith("One line"))
    ans = next(sh for t, sh in boxes.items() if t.startswith("Because it"))
    gap = (ans.top - (copy_box.top + copy_box.height)) / IN
    assert abs(gap - 0.4) < 0.03, gap


def test_kpi_figures_share_one_baseline(tmp_path):
    """Figures of a KPI row sit in boxes of one top and one height (so one baseline), and their labels start together."""
    items = [{"value": "2.9M", "label": "A short label."}, {"value": "3.1%", "label": "A much longer label that wraps over three or four lines in its card."},
             {"value": "31 to 9 min", "label": "Medium label here."}]
    r, prs = _render(tmp_path, _slide_deck("traction", [{"kind": "stats", "items": items}], section="Traction"))
    assert r.layouts["sx"] == "kpi_row"
    figs = [sh for sh in prs.slides[0].shapes if sh.has_text_frame and sh.text_frame.text.strip() in ("2.9M", "3.1%", "31 to 9 min")]
    assert len(figs) == 3
    assert len({(sh.top, sh.height) for sh in figs}) == 1
    labs = [sh for sh in prs.slides[0].shapes if sh.has_text_frame and sh.text_frame.text.strip().endswith(("label.", "here.", "card."))]
    assert len({sh.top for sh in labs}) == 1


def _last_two_words_share_a_line(r, text, width, size, bold=True, font=None):
    """True when the renderer's own wrap estimate does not leave a lone word on the last line of ``text``."""
    lines = r.wrap(R.glue(text), width, size, bold, font or r.th.head)
    return len(lines) == 1 or sum(len(w.split(R.NBSP)) for w in lines[-1]) >= 2   # a glued pair counts as two words


def test_no_title_ends_on_a_lone_word_in_the_golden_deck(tmp_path):
    """Every drawn title (the renderer's wrap estimate at its own size and width) has at least two words on its last line."""
    deck = _deck()
    r, prs = _render(tmp_path, deck)
    n = 0
    for sl in prs.slides:
        for sh in sl.shapes:
            if sh.name == "Title" and sh.has_text_frame:
                run = sh.text_frame.paragraphs[0].runs[0]
                assert _last_two_words_share_a_line(r, run.text.replace(R.NBSP, " "), sh.width / IN, run.font.size.pt), run.text
                n += 1
    assert n >= 10


def test_no_card_text_ends_on_a_lone_word(tmp_path):
    """Card sentences wrap with two words on the last line (the no-break space), whatever the width."""
    r, prs = _render(tmp_path, _slide_deck("insight", [{"kind": "bullets", "items": [
        {"text": "An awareness objective implies most of the audience does not yet know the brand."},
        {"text": "Sparkling iced tea may need a reason to pick it up for the first time."},
        {"text": "With no budget given, reach must be earned through ideas worth sharing as well as bought."}]}]))
    for sh in prs.slides[0].shapes:
        if sh.has_text_frame and len(sh.text_frame.text) > 40:
            run = sh.text_frame.paragraphs[0].runs[0]
            assert _last_two_words_share_a_line(r, run.text.replace(R.NBSP, " "), sh.width / IN, run.font.size.pt, bold=False, font=r.th.body), run.text


def test_glue_changes_no_word():
    """The no-break space joins the last two words and nothing else: the words and their order are unchanged."""
    for t in ["One two three four.", "Two words", "A single", "Many words in a sentence that goes on."]:
        assert norm(R.glue(t)) == norm(t)
    assert R.glue("one two three") == "one two" + R.NBSP + "three"
    assert R.glue("one two") == "one two" and R.glue("one") == "one"
    assert R.glue("a b c\nd e f") == "a b c\nd e f"   # a text with line breaks is left alone


def test_a_short_last_row_spans_the_width_no_empty_cell(tmp_path):
    """Five, four and seven short cards: every row fills the width of the content area (no empty cell)."""
    for n in (4, 5, 7):
        items = [{"text": f"Claim number {i} reads like this."} for i in range(n)]
        r, prs = _render(tmp_path, _slide_deck("insight", [{"kind": "bullets", "items": items}]))
        cards = _cards(prs.slides[0], r.th)
        rows = {}
        for c in cards:
            rows.setdefault(round(c[1], 2), []).append(c)
        assert len(cards) == n
        for row in rows.values():
            left = min(c[0] for c in row)
            right = max(c[0] + c[2] for c in row)
            assert abs(left - r.x0) < 0.02 and abs(right - r.x1) < 0.02, (n, row)


def test_role_grid_has_both_cells_in_every_row(tmp_path):
    """A role grid row is a dark role cell and a panel at one top and one height: no row has an empty cell."""
    steps = [["Strategy lead", "Sets the plan."], ["Creative lead", "Makes the work."], ["Media lead", "Buys the space."]]
    r, prs = _render(tmp_path, _slide_deck("founders", [{"kind": "steps", "items": steps}], section="Team"))
    assert r.layouts["sx"] == "role_grid"
    sl = prs.slides[0]
    dark = [x for x in _fills(sl, r.th.ink) if x[2] > 2]
    light = [x for x in _fills(sl, r.th.panel) if x[2] > 4]
    assert len(dark) == len(light) == 3
    for d, l in zip(dark, light):
        assert abs(d[1] - l[1]) < 0.01 and abs(d[3] - l[3]) < 0.01
    assert abs(max(x[1] + x[3] for x in light) - r.th.grid.body_end) < 0.08


def test_table_width_fits_its_content_and_is_never_a_sixty_percent_stub(tmp_path):
    """A native table is at least as wide as its content estimate; a wide table is not squeezed beside a bullet."""
    rows = [["Alternative", "What it does", "What it misses"],
            ["Paper dockets and a clerk", "Chases disputes after complaints", "25-40 minutes per dispute; credits usually issued"],
            ["Three ERPs most customers use", "Record the invoice", "Do not record proof of delivery"]]
    blocks = [{"kind": "bullets", "items": [{"text": "One line of context."}]}, {"kind": "table", "rows": rows}]
    r, prs = _render(tmp_path, _slide_deck("competition", blocks, section="Competition"))
    tbl = next(sh for sh in prs.slides[0].shapes if getattr(sh, "has_table", False) and sh.has_table)
    size = tbl.table.cell(1, 0).text_frame.paragraphs[0].runs[0].font.size.pt
    assert tbl.width / IN >= r._table_nat(rows, size) - 0.05
    assert tbl.width / IN >= 0.95 * (r.x1 - r.x0)
    # a narrow table beside a bullet is as wide as its content (not the old fixed 60%) and the aside is as tall as it
    nar = [["A", "B"], ["x", "y"], ["z", "w"]]
    r2, prs2 = _render(tmp_path, _slide_deck("competition", [{"kind": "table", "rows": nar}, {"kind": "bullets", "items": [{"text": "Aside."}]}], section="Competition"))
    t2 = next(sh for sh in prs2.slides[0].shapes if getattr(sh, "has_table", False) and sh.has_table)
    assert t2.width / IN >= r2._table_nat(nar, 14) - 0.05
    assert t2.width / IN <= 0.62 * (r2.x1 - r2.x0) + 0.01


def test_chart_has_one_scale_and_does_not_repeat_its_bars_as_number_panels(tmp_path):
    """The chart slide holds one native chart on a 0 to 100 axis with the hero bars in the accent, and no big-number
    panel repeating the figures the bars show."""
    inp = Inputs(brief={}, facts={"f_a": {"value": 58, "unit": "%", "label": "a"}, "f_b": {"value": 47, "unit": "%", "label": "b"},
                                  "f_c": {"value": 62, "unit": "%", "label": "c"}})
    ch = {"kind": "chart", "chart": {"type": "bar", "series": [{"label": "Slump drinkers", "ref": "f_a"}, {"label": "Low sugar buyers", "ref": "f_b"},
                                                                {"label": "Prefer local", "ref": "f_c"}], "hero": ["Slump drinkers", "Prefer local"]}}
    deck = _slide_deck("chart", [ch, {"kind": "text", "text": "The hour and the habit are already mass-market."}], section="Insight")
    out = tmp_path / "c.pptx"
    r = R.Renderer(deck, inp, theme="napkin")
    r.render(out)
    sl = Presentation(str(out)).slides[0]
    assert r.layouts["sx"] == "chart"
    charts = [sh for sh in sl.shapes if sh.has_chart]
    assert len(charts) == 1
    va = charts[0].chart.value_axis
    assert va.minimum_scale == 0 and va.maximum_scale == 100
    xml = charts[0].chart._chartSpace.xml
    assert xml.count(r.th.accent) >= 2 and "<c:dPt>" in xml   # the hero bars carry the accent
    text = slide_text(sl)
    assert "58%" not in text and "62%" not in text


def test_floors_agree_between_theme_and_stage_h():
    """The theme records the floors once and qa_render reads them: body 12, caption 10, footer 9, title 28."""
    S = T.Scale()
    assert qa_render.FLOORS == {"title": float(S.title_min), "body": float(S.body_floor), "caption": float(S.caption_floor), "foot": float(S.foot_floor)}
    assert (S.title_min, S.body_floor, S.caption_floor, S.foot_floor) == (28, 12, 10, 9)
    assert S.foot_floor <= S.caption_floor <= S.caption <= S.body_floor <= S.body
    assert S.foot >= S.foot_floor and S.min <= S.caption_floor
    assert T.check_palette(T.NAPKIN) == []
    assert any("caption" in p for p in T.check_palette(T.NAPKIN.__class__(**{**T.NAPKIN.__dict__, "scale": T.Scale(caption=9)})))
    assert {"Watermark", "SlideNumber", "Footnotes", "Section", "CheckFailed"} <= qa_render.CHROME


def test_no_min_font_findings_on_slides_that_fit_at_the_floors(tmp_path):
    """On the golden deck stage h finds no min_font on any slide that was not set at the last-resort sizes."""
    deck = _deck()
    out = tmp_path / "g.pptx"
    r = R.Renderer(deck, _inp(), theme="napkin")
    r.render(out)
    res = qa_render.check_pptx(out)
    for n, fs in res["slides"].items():
        for f in fs:
            if f["check"] == "min_font":
                assert r.layouts[deck["order"][int(n) - 1]] in ("placeholder_card", "flow") or "title" in f["message"], (n, f)
    assert res["summary"]["by_check"].get("min_font", 0) <= 40


def test_relaxed_last_resort_is_only_for_a_slide_that_cannot_fit_at_the_floors(tmp_path):
    """A normal slide has no run under 12 pt except captions (10 pt+), footer lines and chips; the relaxed flag stays off."""
    r, prs = _render(tmp_path, _slide_deck("insight", SHORT))
    assert r._relaxed is False
    for sh in prs.slides[0].shapes:
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                for run in p.runs:
                    if run.text.strip():
                        floor = 9 if sh.name in qa_render.CHROME else (10 if sh.name == "Caption" else 12)
                        assert run.font.size.pt >= floor, (sh.name, run.font.size.pt)


def test_long_titles_stay_readable_and_every_word_survives(tmp_path):
    """A 90-character title is set at 28 pt or more, balanced over its lines, and every word is on the slide."""
    title = "Most young Irish adults already want what the brand is but they do not yet know its name"
    deck = _slide_deck("insight", SHORT, title=title)
    r, prs = _render(tmp_path, deck)
    t = next(sh for sh in prs.slides[0].shapes if sh.name == "Title")
    assert t.text_frame.paragraphs[0].runs[0].font.size.pt >= 28
    assert norm(title) in slide_text(prs.slides[0])


def test_control_characters_in_deck_strings_do_not_abort_the_build(tmp_path):
    """NUL, VT and US in deck text are stripped; the deck still renders."""
    d = json.loads(json.dumps(_deck()).replace("Lumen", "Lu\\u0000m\\u000ben\\u001fX"))
    out = R.Renderer(d, _inp(), theme="napkin").render(tmp_path / "x.pptx")
    assert Presentation(str(out)).slides


def test_non_string_brand_fonts_are_ignored():
    """Fonts that are not plain strings never reach the theme."""
    t = T.get_theme("napkin", {"head": ["x"], "body": 123})
    assert t.head == T.NAPKIN.head and t.body == T.NAPKIN.body
    assert T.get_theme("napkin", {"head": "Georgia"}).head == "Georgia"


def test_non_image_logo_path_degrades_to_no_logo(tmp_path):
    """A logo path that is not an image file (e.g. a text file) is skipped, not opened as a picture."""
    d = _deck()
    f = tmp_path / "x.png"
    f.write_text("not an image")
    d["meta"]["brand"] = {"logo_dark_png": str(f)}
    out = R.Renderer(d, _inp(), theme="napkin").render(tmp_path / "y.pptx")
    assert Presentation(str(out)).slides
