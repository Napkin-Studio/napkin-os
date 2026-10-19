"""Offline tests for the per-type polish (roadmap items 9, 10 and 11): the tender's per-criterion answers and weight share, the agency
risks, competition and measurement-plan slides, the investor's required shape and the layouts of the new slide types. No model call: a stage
that would call a model is tested with a stubbed ``_call`` and the recorded JSON fixtures under fixtures/polish/ (every client invented).
Run from ``engine/``: ``PYTHONPATH=. python3 -m pytest -q deck/tests``.
"""
from __future__ import annotations

import json
import re
import shutil
from pathlib import Path

import pytest
from pptx import Presentation

from deck import assemble, inputs, profiles, render_pptx as R, run_deck, spec, stages, validate
from deck.spec import shown_blocks

HERE = Path(__file__).parent
FIX = HERE / "fixtures"
LUMEN, POLISH, TYPES = FIX / "lumen_water", FIX / "polish", FIX / "types"
CRITERIA = ("Team and personnel", "Insight and campaign strategy", "Measurement and reporting")


# ---------------------------------------------------------------- helpers

def _lumen(extras: bool = True):
    """(deck, inputs) for the Lumen Water brief; ``extras`` swaps in the stage d output that also writes risks, competition and a measurement plan."""
    st = {k: json.loads((LUMEN / f).read_text()) for k, f in (("a", "a_requirements.json"), ("b", "b_storyline.json"), ("c", "c_territories.json"), ("d", "d_idea.json"))}
    if extras:
        st["d"] = json.loads((POLISH / "d_idea_extras.json").read_text())
    inp = inputs.load(LUMEN / "brief_object.json")
    return assemble.build(inp, st, client="Lumen Water", project="National awareness campaign"), inp


def _types(deck):
    """Slide types in deck order."""
    return [deck["slides"][s]["type"] for s in deck["order"]]


def _of(deck, t):
    """The first slide of a type."""
    return next(s for s in deck["slides"].values() if s["type"] == t)


def _tender_stages(tmp_path, with_answers=True) -> Path:
    """A stage directory: the tender fixture's recorded stages, and (``with_answers``) the three recorded per-criterion answers."""
    d = tmp_path / "stages"
    shutil.copytree(TYPES / "tender_response", d, ignore=shutil.ignore_patterns("prompt.txt"))
    if with_answers:
        for f in (POLISH / "tender").glob("d_*.json"):
            shutil.copy(f, d / f.name)
    return d


def _tender(tmp_path, with_answers=True):
    """(deck, inputs) for the tender fixture built through run_deck's own step (_criterion_answers) from recorded files."""
    sd = _tender_stages(tmp_path, with_answers)
    inp = inputs.load(None, prompt_text=(TYPES / "tender_response" / "prompt.txt").read_text().strip())
    prof = profiles.load("tender_response")
    st = {k: json.loads((sd / f).read_text()) for k, f in (("a", "a_requirements.json"), ("b", "b_storyline.json"), ("c", "c_territories.json"), ("d", "d_idea.json"))}
    args = type("A", (), {"stages": sd, "live": False})()
    rep = {"stages": {}}
    run_deck._criterion_answers(args, inp, st, prof, rep)
    return assemble.build(inp, st, client="Harbour", project="Recycling", profile=prof), inp, rep, st


# ---------------------------------------------------------------- item 10: agency risks, competition, measurement plan

def test_the_agency_deck_presents_risks_before_the_investment_and_never_last():
    """Risks are presented (not in the appendix), placed before the investment and the ask, and the presented story does not end on them."""
    deck, _ = _lumen()
    pres = [deck["slides"][s]["type"] for s in validate.presented(deck)]
    assert "risks" in pres and pres.index("risks") < pres.index("investment") and pres[-1] != "risks"
    rk = _of(deck, "risks")
    kv = rk["blocks"][0]
    assert kv["kind"] == "kv" and [k for k, _ in kv["items"]][0] == "Households read the saving ask as a bill warning"
    assert len(kv["items"]) == 3 and rk["status"] == "drafted"


def test_the_measurement_plan_extends_measurement_with_baseline_target_lag_and_confounders():
    """A table of KPI, baseline, target, lag and confounders, right after the measurement slide; a baseline the inputs do not give is said so."""
    deck, _ = _lumen()
    t = _types(deck)
    assert t.index("measurement_plan") == t.index("measurement") + 1
    rows = _of(deck, "measurement_plan")["blocks"][0]["rows"]
    assert rows[0] == ["KPI", "Baseline", "Target", "Lag", "Confounders"] and rows[1][1] == "None in the brief: set before launch"
    assert rows[1][4].startswith("Weather and a tariff change")


def test_competition_names_only_alternatives_the_inputs_cite():
    """An alternative with no ref is dropped (the model may not name a rival the inputs do not); the plan to win follows when it fits."""
    deck, _ = _lumen()
    c = _of(deck, "competition")
    rows = c["blocks"][0]["rows"]
    assert [r[0] for r in rows[1:]] == ["Keep the current bill inserts"] and c["blocks"][0]["compare"] is True
    assert any(b["kind"] == "callout" for b in c["blocks"]) and c["blocks"][0]["refs"] == ["brief:background"]


def test_without_recorded_alternatives_competition_and_risks_are_labelled_placeholders():
    """A stage d that wrote none of them leaves a labelled placeholder asking the person: the code writes no risk and names no rival itself."""
    deck, inp = _lumen(extras=False)
    for t in ("competition", "risks"):
        sl = _of(deck, t)
        assert sl["status"] == "needs_human"
        item = sl["blocks"][0]["items"][0]
        assert item["human"] is True and item["text"].startswith(spec.PLACEHOLDER_PREFIX) and "To confirm:" in item["text"]
    assert "measurement_plan" not in _types(deck)   # no plan recorded: no slide
    assert not [f for f in validate.run(deck, inp) if f["level"] == "error"]


def test_new_slide_types_have_budgets_and_schema_entries():
    """The budgets live in spec and the schema's enum lists the same types; the plan's word budget for the measurement plan is 110."""
    assert spec.SLIDE_TYPES["measurement_plan"]["words"] == 110 and spec.SLIDE_TYPES["risks"]["words"] == 90 and spec.SLIDE_TYPES["competition"]["words"] == 80
    schema = json.loads((Path(R.__file__).parent / "deck.schema.json").read_text())
    assert schema["definitions"]["slide"]["properties"]["type"]["enum"] == list(spec.SLIDE_TYPES)
    deck, inp = _lumen()
    for t in ("risks", "competition", "measurement_plan"):
        sl = _of(deck, t)
        assert validate.words(sl) <= spec.SLIDE_TYPES[t]["words"], t
    assert not [f for f in validate.run(deck, inp) if f["check"] == "word_budget"]


def test_agency_spine_places_the_new_types():
    """agency_pitch's spine lists competition after the proposition, the plan after measurement and risks before the investment."""
    sp = profiles.load("agency_pitch")["spine"]
    assert sp.index("competition") == sp.index("proposition") + 1 and sp.index("measurement_plan") == sp.index("measurement") + 1
    assert sp.index("risks") < sp.index("investment")


def test_a_stated_limit_trims_the_new_slides_last():
    """Held to a small limit the story gives up competition, then the plan and risks, only after the other trimmable types."""
    deck, _ = _lumen()
    story = [deck["slides"][s] for s in validate.presented(deck)]
    story_ = list(story)
    out, back = assemble._trim(list(deck["slides"].values()), story_, [], 12)
    kept = [s["type"] for s in out]
    assert "risks" in kept and "competition" in kept and len(out) == 12
    out8, _ = assemble._trim(list(deck["slides"].values()), list(story), [], 8)
    assert len(out8) == 8 and "competition" not in [s["type"] for s in out8]


def _layouts(tmp_path, deck):
    """Render a deck dict and return (renderer, presentation)."""
    from deck.inputs import Inputs
    r = R.Renderer(deck, Inputs(brief={}), theme="napkin")
    out = tmp_path / "d.pptx"
    r.render(out)
    return r, Presentation(str(out))


def _text(sl):
    """Whitespace-normalised text of a rendered slide (text frames and table cells)."""
    out = []
    for sh in sl.shapes:
        if sh.has_text_frame:
            out.append(" ".join(p.text for p in sh.text_frame.paragraphs))
        elif getattr(sh, "has_table", False) and sh.has_table:
            out += [c.text for row in sh.table.rows for c in row.cells]
    return re.sub(r"\s+", " ", " ".join(out)).strip()


def test_each_new_type_has_its_layout_and_keeps_every_word(tmp_path):
    """Risk cards, a comparison grid, a native table and the objective's big-number panel, none fallen back to the flow, and no word lost."""
    deck, _ = _lumen()
    r, prs = _layouts(tmp_path, deck)
    want = {"risks": "risk_cards", "competition": "comparison_grid", "measurement_plan": "native_table", "objective": "objective_panel"}
    for sid, slide in zip(deck["order"], prs.slides):
        s = deck["slides"][sid]
        if s["type"] in want:
            assert R.choose_layout(s) == want[s["type"]] and r.layouts[sid] == want[s["type"]], (s["type"], r.layouts[sid])
            text = _text(slide)
            words = [s["title"]]
            for b in shown_blocks(s):
                if b["kind"] == "kv":
                    words += [x for k, v in b["items"] for x in (k, v)]
                elif b["kind"] == "table":
                    words += [c for row in b["rows"] for c in row]
                elif b["kind"] == "bullets":
                    words += [i["text"] for i in b["items"]]
                elif "text" in b:
                    words.append(b["text"])
            for w in words:
                assert re.sub(r"\s+", " ", w).strip() in text, (s["type"], w)


def test_placeholder_risks_and_competition_render_as_placeholder_cards(tmp_path):
    """The labelled placeholder is never styled as finished: a deck with no recorded risks draws dashed placeholder cards."""
    deck, _ = _lumen(extras=False)
    r, _ = _layouts(tmp_path, deck)
    for sid, s in deck["slides"].items():
        if s["type"] in ("risks", "competition"):
            assert r.layouts[sid] == "placeholder_card"


def test_objective_facts_pull_target_baseline_and_date_from_the_sentence():
    """The panel's figures are the sentence's own: target after 'to' in a from/to, baseline after 'from', date from 'by' or 'in N months'."""
    assert R._objective_facts(["Cut household use by 5% in 12 months"]) == {"target": "5%", "baseline": "", "date": "in 12 months"}
    assert R._objective_facts(["Move awareness from 12% to 30% by Q4 2027"]) == {"target": "30%", "baseline": "12%", "date": "by Q4 2027"}
    assert R._objective_facts(["Raise EUR 450,000 by June"])["target"] == "EUR 450,000"
    assert R._objective_facts(["People see water as worth looking after"]) is None


def test_an_objective_without_a_figure_is_drawn_as_cards():
    """No figure to pull out: the objective slide keeps the cards layout (not the panel, not the flow)."""
    sl = {"type": "objective", "status": "from_source", "title": "T", "claims": [], "blocks": [{"kind": "bullets", "items": [{"text": "People see water as worth looking after"}]}]}
    assert R.choose_layout(sl) == "cards"
    sl["blocks"][0]["items"][0]["text"] = "Cut use by 5% in 12 months"
    assert R.choose_layout(sl) == "objective_panel"


# ---------------------------------------------------------------- item 11: investor

def test_investor_story_puts_the_objective_second_and_what_we_know_after_traction():
    """The objective (round and milestones) is slide 2, why_now is present, what_we_know follows traction; the cover comes first."""
    inp = inputs.load(None, prompt_text=(POLISH / "investor" / "prompt.txt").read_text().strip())
    st = {"i": json.loads((POLISH / "investor" / "i_investor.json").read_text())}
    deck = assemble.build_investor(inp, st, client="Fenwick", project="Seed", profile=profiles.load("investor_seed"))
    t = _types(deck)
    assert t[:3] == ["title", "objective", "problem"] and "why_now" in t
    assert t.index("what_we_know") == t.index("traction") + 1 and t.index("objective") == 1
    assert _of(deck, "objective")["blocks"][0]["kind"] == "stats" and _of(deck, "what_we_know")["section"] == "What we know"
    assert not [f for f in validate.run(deck, inp) if f["check"] in ("word_budget", "numbers_trace") and f["level"] == "error"]


def test_investor_spine_and_prompt_carry_the_required_shape():
    """The spine lists objective second and what_we_know after traction; the i prompt asks for why_now, the round's objective and the scariest number."""
    sp = profiles.load("investor_seed")["spine"]
    assert sp[:2] == ["title", "objective"] and sp.index("what_we_know") == sp.index("traction") + 1
    assert {"objective", "what_we_know"} <= set(spec.INVESTOR_TYPES)
    text = (Path(stages.__file__).parent / "prompts" / "i_investor.md").read_text()
    for must in ("why_now` is required", "what_we_know", "scariest number", "milestones the money buys"):
        assert must in text
    sysmsg = stages._sysmsg("i_investor.md", "i", profiles.load("investor_seed"))
    assert "<type_guidance>" in sysmsg and "why_now is required" in sysmsg and sysmsg.index("<type_guidance>") < sysmsg.index("Inputs are the only source:")


def test_investor_golden_is_unchanged_by_the_new_types():
    """The recorded investor golden has none of the new sections, so it still builds byte for byte (the spine change reorders nothing there)."""
    inp = inputs.load(None, prompt_text=(TYPES / "investor_seed" / "prompt.txt").read_text().strip())
    st = {"i": json.loads((TYPES / "investor_seed" / "i_investor.json").read_text())}
    deck = assemble.build_investor(inp, st, client="Fixture Client", project="Fixture Project", profile=profiles.load("investor_seed"))
    deck["meta"].pop("built_at"); deck["meta"].pop("built_from")
    assert json.loads(json.dumps(deck, default=str)) == json.loads((TYPES / "investor_seed" / "golden_deck.json").read_text())


# ---------------------------------------------------------------- item 9: tender per-criterion answers

def test_budgets_come_from_weights_and_the_heaviest_gets_the_most_words():
    """30/40/30 gives the 40% criterion the most; a page limit caps; a missing weight takes the average; every budget is at least the floor."""
    scoring = [{"criterion": "A", "weight": "30%", "page_limit": "4 pages"}, {"criterion": "B", "weight": "40%", "page_limit": "6 pages"}, {"criterion": "C", "weight": "30%"}]
    b = assemble.criterion_budgets(scoring)
    assert b == {"A": 180, "B": 240, "C": 180} and max(b, key=b.get) == "B"
    assert assemble.criterion_budgets([{"criterion": "A", "weight": "10%", "page_limit": "0.1 pages"}, {"criterion": "B", "weight": "90%"}])["A"] == 45
    assert assemble.criterion_budgets([{"criterion": "A", "weight": "1%"}, {"criterion": "B", "weight": "99%"}])["A"] == spec.ANSWER_WORDS_MIN
    assert assemble.criterion_budgets([{"criterion": "A", "weight": "50%", "page_limit": "60 words"}, {"criterion": "B", "weight": ""}])["A"] == 60
    assert assemble.criterion_budgets([{"criterion": "A"}, {"criterion": "B"}]) == {"A": 300, "B": 300}


def test_the_stage_makes_one_call_per_criterion_with_its_budget(monkeypatch):
    """A stubbed ``_call``: one call each, the budget in the user turn, the criterion named, the output recorded with its budget; no model runs."""
    calls = []
    monkeypatch.setattr(stages, "_call", lambda stage, user, system: (calls.append((stage, user, system)) or ({"title": "T", "points": []}, {"model": "stub", "calls": 1})))
    inp = inputs.load(None, prompt_text="x")
    scoring = json.loads((TYPES / "tender_response" / "a_requirements.json").read_text())["scoring"]
    budgets = assemble.criterion_budgets(scoring)
    for c in scoring:
        obj, _ = stages.criterion_answer(inp, {"a": {"requirements": []}}, c, budgets[c["criterion"]], profiles.load("tender_response"))
        assert obj["criterion"] == c["criterion"] and obj["word_budget"] == budgets[c["criterion"]]
    assert [c[0] for c in calls] == ["d_criterion"] * 3
    assert all(f"<word_budget>{budgets[c['criterion']]}</word_budget>" in u for (_, u, _), c in zip(calls, scoring))
    assert "<type_guidance>" in calls[0][2] and "Do only the work the requirements ask for" in calls[0][2]   # the tender's stage d guidance reaches the call
    assert "d_criterion" in stages.MODELS and "d_criterion" in stages.EFFORT


def test_recorded_answers_become_the_deck_slides_in_the_buyers_order(tmp_path):
    """Each recorded criterion answer is a criterion_answer slide after the objective, in the scoring order, carrying its criterion and budget."""
    deck, inp, rep, st = _tender(tmp_path)
    types = _types(deck)
    assert types.index("criterion_answer") == types.index("objective") + 1
    ans = [s for s in deck["slides"].values() if s["type"] == "criterion_answer"]
    assert [s["criterion"] for s in ans] == list(CRITERIA) and [s["budget"] for s in ans] == [180, 240, 180]
    assert rep["stages"]["d_criteria"]["source"] == "recorded" and rep["stages"]["d_criteria"]["answered"] == 3
    assert deck["meta"]["stages"]["d_criteria"].startswith("recorded stand-in")
    team = ans[0]
    assert team["status"] == "needs_human" and team["blocks"][0]["items"][0]["human"] is True and "team CVs" in team["blocks"][0]["items"][0]["text"]
    mid = ans[1]
    assert mid["status"] in ("from_source", "drafted") and len(mid["blocks"][0]["items"]) == 6 and validate.words(mid) <= mid["budget"]


def test_the_heaviest_criterion_holds_the_most_words_and_every_answer_is_within_its_budget(tmp_path):
    """The measure of item 9: every answer is within its budget (so its page limit) and the 40% criterion holds the most words."""
    deck, inp, _, _ = _tender(tmp_path)
    ans = {s["criterion"]: s for s in deck["slides"].values() if s["type"] == "criterion_answer"}
    w = {k: validate.words(v) for k, v in ans.items()}
    assert all(w[k] <= ans[k]["budget"] for k in ans)
    assert max(w, key=w.get) == "Insight and campaign strategy"
    assert not [f for f in validate.run(deck, inp) if f["check"] == "word_budget"]


def test_a_point_that_does_not_fit_is_kept_in_the_notes_never_cut(tmp_path):
    """A budget smaller than the answer shows whole points only and names the rest in the notes: nothing is reworded or silently dropped."""
    sc = [{"criterion": "C1", "weight": "50%", "page_limit": "60 words"}, {"criterion": "C2", "weight": "50%"}]
    pts = [{"text": "word " * 25, "refs": ["prompt:1"]} for _ in range(4)]
    st = {"d_criteria": {"answers": [{"criterion": "C1", "title": "T one for the first", "points": pts}]}}
    (sl,) = assemble._criterion_slides(st, sc)
    assert len(sl["blocks"][0]["items"]) == 2 and "2 further point(s) did not fit" in sl["notes"] and sl["budget"] == 60
    st["d_criteria"]["answers"][0]["points"] = [{"text": "word " * 90, "refs": ["prompt:1"]}]
    (sl,) = assemble._criterion_slides(st, sc)
    assert sl["status"] == "needs_human" and "did not fit" not in sl["notes"]


def test_without_recorded_answers_the_stage_is_skipped_quietly(tmp_path):
    """No d_<criterion>.json and no --live: skipped, with a report line, no slide, and the deck builds exactly as it did before."""
    deck, _, rep, st = _tender(tmp_path, with_answers=False)
    assert rep["stages"]["d_criteria"]["source"] == "skipped: no recorded output" and "d_criteria" not in st
    assert "criterion_answer" not in _types(deck) and "d_criteria" not in deck["meta"]["stages"]


def test_live_calls_each_criterion_once_and_records_its_file(tmp_path, monkeypatch):
    """With --live a missing answer is made by stages.criterion_answer (stubbed here), recorded as d_<criterion>.json and counted in the stage's cost."""
    sd = _tender_stages(tmp_path, with_answers=False)
    seen = []
    def fake(inp, st, c, budget, profile=None):
        """A stub for the model call: records the criterion and budget it was given."""
        seen.append((c["criterion"], budget))
        return {"stage": "d_criterion", "criterion": c["criterion"], "title": "A takeaway for the criterion", "points": [{"text": "A cited point", "refs": ["prompt:2"]}],
                "produced_by": "stub"}, {"model": "stub", "calls": 1, "prompt_tokens": 10, "completion_tokens": 5, "secs": 0, "usd": 0.01}
    monkeypatch.setattr(stages, "criterion_answer", fake)
    inp = inputs.load(None, prompt_text=(TYPES / "tender_response" / "prompt.txt").read_text().strip())
    st = {"a": json.loads((sd / "a_requirements.json").read_text())}
    rep = {"stages": {}}
    run_deck._criterion_answers(type("A", (), {"stages": sd, "live": True})(), inp, st, profiles.load("tender_response"), rep)
    assert seen == [("Team and personnel", 180), ("Insight and campaign strategy", 240), ("Measurement and reporting", 180)]
    assert sorted(p.name for p in sd.glob("d_*.json")) == ["d_idea.json", "d_insight_and_campaign_strategy.json", "d_measurement_and_reporting.json", "d_team_and_personnel.json"]
    row = rep["stages"]["d_criteria"]
    assert row["source"] == "live" and row["calls"] == 3 and abs(row["usd"] - 0.03) < 1e-9
    # a second, recorded run reads the files back and makes no call
    seen.clear()
    rep2 = {"stages": {}}
    run_deck._criterion_answers(type("A", (), {"stages": sd, "live": False})(), inp, {"a": st["a"]}, profiles.load("tender_response"), rep2)
    assert not seen and rep2["stages"]["d_criteria"]["source"] == "recorded"


def test_other_types_run_no_criterion_step(tmp_path):
    """Only a type whose spine holds criterion_answer (the tender response) runs the step; the agency pitch adds nothing to the report."""
    inp = inputs.load(None, prompt_text="x")
    rep = {"stages": {}}
    st = {"a": {"scoring": [{"criterion": "A", "weight": "10%"}]}}
    run_deck._criterion_answers(type("A", (), {"stages": tmp_path, "live": True})(), inp, st, profiles.load("agency_pitch"), rep)
    assert rep["stages"] == {} and "d_criteria" not in st


def test_weight_share_warns_for_a_heavy_criterion_with_under_half_its_share_and_reports_the_minimum(tmp_path):
    """The team criterion (30%) has no words (a placeholder): weight_share warns on it and names the lowest ratio; the report carries the minimum."""
    deck, inp, _, _ = _tender(tmp_path)
    rows = {r["criterion"]: r for r in validate.weight_share_rows(deck)}
    assert rows["Team and personnel"]["ratio"] < spec.WEIGHT_SHARE_FLOOR and rows["Team and personnel"]["heavy"] is True
    assert abs(sum(r["weight_share"] for r in rows.values()) - 1.0) < 1e-3
    f = [x for x in validate.run(deck, inp) if x["check"] == "weight_share"]
    assert f and all(x["level"] == "warn" for x in f) and "Team and personnel" in f[0]["message"] and "lowest ratio" in f[0]["message"]
    rep = validate.weight_share_report(deck)
    assert rep["min_criterion"] in rows and rep["min_ratio"] == min(r["ratio"] for r in rows.values())


def test_weight_share_is_quiet_when_the_words_follow_the_weights(tmp_path):
    """A deck whose answers use their budgets in proportion to their weights raises no weight_share warning."""
    deck, inp, _, _ = _tender(tmp_path)
    for s in deck["slides"].values():   # words in the proportion of the weights: 30, 40, 30
        if s["type"] == "criterion_answer":
            n = 4 if s["criterion"].startswith("Insight") else 3
            s["status"] = "from_source"
            s["blocks"] = [{"kind": "bullets", "items": [{"text": "A named lead with ten years on public campaigns", "refs": ["prompt:2"]} for _ in range(n)]}]
    assert not [x for x in validate.run(deck, inp) if x["check"] == "weight_share"]


def test_response_map_shows_weight_and_the_share_of_words(tmp_path):
    """The scoring slide's table gains a 'Share of our words' column beside the weight; the answer slides are the slides it points to."""
    deck, _, _, _ = _tender(tmp_path)
    rows = _of(deck, "scoring")["blocks"][0]["rows"]
    assert rows[0] == ["What you score", "Weight", "Share of our words", "Where we answer it"]
    assert [r[1] for r in rows[1:]] == ["30%", "40%", "30%"] and all(re.fullmatch(r"\d+%", r[2]) for r in rows[1:])
    num = {s: i for i, s in enumerate(deck["order"], 1)}
    ans = [num[s] for s, v in deck["slides"].items() if v["type"] == "criterion_answer"]
    for r, n in zip(rows[1:], ans):
        assert str(n) in r[3].split(", ")


def test_a_criterion_with_a_written_answer_is_not_reported_unanswered(tmp_path):
    """criteria_map counts the answer slide as answering the whole criterion (the concept words of its parts no longer have to match a slide type)."""
    deck, _, _, _ = _tender(tmp_path)
    assert all(not c["missing"] for c in assemble.criteria_coverage(deck))


def test_run_deck_writes_the_weight_share_report(tmp_path, monkeypatch):
    """End to end with the renderers stubbed: report.json carries the d_criteria stage and the weight_share minimum; validation.json the warning."""
    sd = _tender_stages(tmp_path)
    for mod, fn in ((run_deck.render_pptx, "render"), (run_deck.render_web, "render")):
        monkeypatch.setattr(mod, fn, lambda *a, **k: tmp_path / "deck.pptx")
    monkeypatch.setattr(run_deck.subprocess, "run", lambda *a, **k: None)
    monkeypatch.setattr(run_deck, "_soffice", lambda: "soffice")
    out = tmp_path / "out"
    run_deck.main(["--prompt-file", str(TYPES / "tender_response" / "prompt.txt"), "--kind", "tender", "--stages", str(sd), "--client", "Harbour", "--project", "Recycling",
                   "--out", str(out), "--no-qa"])
    rep = json.loads((out / "report.json").read_text())
    assert rep["stages"]["d_criteria"]["answered"] == 3 and rep["weight_share"]["min_criterion"] == "Team and personnel"
    assert any(f["check"] == "weight_share" for f in json.loads((out / "validation.json").read_text()))


def test_every_recorded_stage_directory_still_builds_without_the_new_stages(tmp_path):
    """The older tender fixture (no d_<criterion> files) runs its steps with the new one skipped quietly and a deck with no answer slides."""
    sd = _tender_stages(tmp_path, with_answers=False)
    rep_out = tmp_path / "o"
    from deck import run_deck as rd
    import types as _t
    rep = {"stages": {}}
    inp = inputs.load(None, prompt_text=(TYPES / "tender_response" / "prompt.txt").read_text().strip())
    rd._criterion_answers(_t.SimpleNamespace(stages=sd, live=False), inp, {"a": json.loads((sd / "a_requirements.json").read_text())}, profiles.load("tender_response"), rep)
    assert rep["stages"]["d_criteria"] == {"source": "skipped: no recorded output", "criteria": 3, "answered": 0,
                                           "per_criterion": {c: {"source": "skipped: no recorded output", "budget": b} for c, b in zip(CRITERIA, (180, 240, 180))}}
    assert rep_out  # a Path only: nothing is written


def test_a_criterion_with_an_answer_has_its_slide_in_the_answer_and_the_concept_slides_in_the_backup(tmp_path):
    """The deck is derived from the answers: the slides whose types answer an answered criterion (insight, team, measurement here) go behind the divider; the answers stay."""
    deck, _, _, _ = _tender(tmp_path)
    pres = [deck["slides"][s]["type"] for s in validate.presented(deck)]
    assert pres.count("criterion_answer") == 3 and not {"insight", "measurement"} & set(pres)
    back = [s["type"] for s in deck["slides"].values()][len(pres):]
    assert {"insight", "measurement"} <= set(back)   # the team answer is only a placeholder, so the team slide is not replaced by it
    assert {"title", "scoring", "objective", "next_step"} <= set(pres)
    # with no recorded answers the same deck keeps those slides presented, as before
    plain, _, _, _ = _tender(tmp_path / "plain", with_answers=False)
    assert "insight" in [plain["slides"][s]["type"] for s in validate.presented(plain)]


# ---------------------------------------------------------------- polish fixes: untraced figures, needs_human answers, slug collisions

def test_a_figure_in_a_risk_is_traced_like_any_other():
    """A risk or mitigation with a number no input states raises numbers_trace (the risks block no longer carries basis 'proposal')."""
    from deck import validate
    st = {k: json.loads((LUMEN / f).read_text()) for k, f in (("a", "a_requirements.json"), ("b", "b_storyline.json"), ("c", "c_territories.json"), ("d", "d_idea.json"))}
    st["d"] = json.loads((POLISH / "d_idea_extras.json").read_text())
    st["d"]["risks"] = [{"risk": "Budget cut by 37% in March", "mitigation": "We reserve EUR 912,000 and cap at 14 weeks", "refs": [], "basis": "proposal"}]
    inp = inputs.load(LUMEN / "brief_object.json")
    deck = assemble.build(inp, st, client="Lumen Water", project="National awareness campaign")
    sid = next(s for s, x in deck["slides"].items() if x["type"] == "risks")
    msgs = " ".join(f["message"] for f in validate.run(deck, inp) if f["check"] == "numbers_trace" and f["slide"] == sid)
    assert "37" in msgs and "912000" in msgs and "14" in msgs


def test_an_uncited_tender_answer_point_cannot_carry_an_invented_figure(tmp_path):
    """A criterion-answer point with no refs is an opinion, not a proposal: its figures are traced."""
    from deck import validate
    sd = _tender_stages(tmp_path)
    f = sd / "d_insight_and_campaign_strategy.json"
    o = json.loads(f.read_text())
    o["points"][0].update(text="We have won 14 recycling awards and will cut waste by 63%.", refs=[])
    f.write_text(json.dumps(o))
    inp = inputs.load(None, prompt_text=(TYPES / "tender_response" / "prompt.txt").read_text().strip())
    prof = profiles.load("tender_response")
    st = {k: json.loads((sd / n).read_text()) for k, n in (("a", "a_requirements.json"), ("b", "b_storyline.json"), ("c", "c_territories.json"), ("d", "d_idea.json"))}
    run_deck._criterion_answers(type("A", (), {"stages": sd, "live": False})(), inp, st, prof, {"stages": {}})
    deck = assemble.build(inp, st, client="Harbour", project="Recycling", profile=prof)
    sid = next(s for s, x in deck["slides"].items() if x.get("criterion") == "Insight and campaign strategy")
    msgs = " ".join(x["message"] for x in validate.run(deck, inp) if x["check"] == "numbers_trace" and x["slide"] == sid)
    assert "63" in msgs and "14" in msgs


def test_a_needs_human_answer_does_not_push_a_real_slide_behind_the_appendix(tmp_path, monkeypatch):
    """An answer that is only a placeholder (no points) leaves the real slides for that criterion in the presented story."""
    monkeypatch.setattr(assemble, "_trim", lambda S, story, backup, limit: (story, backup))   # the prompt's slide limit is not under test
    sd = _tender_stages(tmp_path)
    (sd / "d_measurement_and_reporting.json").write_text(json.dumps({"stage": "d_criterion", "criterion": "Measurement and reporting", "points": [], "needs": "x"}))
    inp = inputs.load(None, prompt_text=(TYPES / "tender_response" / "prompt.txt").read_text().strip())
    prof = profiles.load("tender_response")
    st = {k: json.loads((sd / n).read_text()) for k, n in (("a", "a_requirements.json"), ("b", "b_storyline.json"), ("c", "c_territories.json"), ("d", "d_idea.json"))}
    run_deck._criterion_answers(type("A", (), {"stages": sd, "live": False})(), inp, st, prof, {"stages": {}})
    deck = assemble.build(inp, st, client="Harbour", project="Recycling", profile=prof)
    from deck import validate
    story = validate.presented(deck)
    assert "measurement" in [deck["slides"][s]["type"] for s in story]


def test_long_criterion_names_sharing_forty_characters_get_distinct_files():
    """Two criteria that differ only after the 40th character map to different stems; short names are unchanged."""
    a = run_deck._slug("Quality of the proposed marketing strategy and approach (Lot A)")
    b = run_deck._slug("Quality of the proposed marketing strategy and approach (Lot B)")
    assert a != b and len(a) <= 40 and len(b) <= 40
    assert run_deck._slug("Team and personnel") == "team_and_personnel"


def test_objective_facts_never_label_a_budget_or_ceiling_the_target():
    """'$5m budget' and 'up to EUR 5m' are limits, not the target: with no other figure the objective is drawn as cards."""
    assert R._objective_facts(["Cut waiting time to 10 days from 2 weeks by Q3 2027, $5m budget"]) is None
    got = R._objective_facts(["Grow donations by 20% within 12 months, to reach new donors up to \u20ac5m"])
    assert got["target"] == "20%"


def test_recorded_criterion_usage_is_summed_and_priced(tmp_path):
    """A fully recorded tender run reports the d_criteria tokens and a priced total (it used to report none)."""
    sd = _tender_stages(tmp_path, True)
    for f in sd.glob("d_*.json"):
        if f.name == "d_idea.json":
            continue
        o = json.loads(f.read_text())
        o["usage"] = {"model": "claude-sonnet-5-5", "prompt_tokens": 1000, "completion_tokens": 500}
        f.write_text(json.dumps(o))
    inp = inputs.load(None, prompt_text=(TYPES / "tender_response" / "prompt.txt").read_text().strip())
    prof = profiles.load("tender_response")
    st = {"a": json.loads((sd / "a_requirements.json").read_text())}
    rep = {"stages": {}}
    run_deck._criterion_answers(type("A", (), {"stages": sd, "live": False})(), inp, st, prof, rep)
    row = rep["stages"]["d_criteria"]
    assert row["source"] == "recorded" and row["prompt_tokens"] == 3000 and row["completion_tokens"] == 1500
    assert run_deck._price_stage(row).get("usd", 0) > 0
