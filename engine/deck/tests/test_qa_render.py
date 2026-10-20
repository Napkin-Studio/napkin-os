"""Stage h: the rendered-slide QA flags each defect, passes a clean slide and never raises."""
from __future__ import annotations

import json
from pathlib import Path

import pytest

pytest.importorskip("pptx")
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from deck import qa_render, run_deck

HERE = Path(__file__).parent
BRIEF = HERE / "fixtures" / "lumen_water"
INK = RGBColor(0x1B, 0x24, 0x30)


def _deck(tmp_path, build):
    """A 13.33 x 7.5 in pptx with one white slide; ``build(slide)`` adds shapes."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    sl.background.fill.solid()
    sl.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
    build(sl)
    p = tmp_path / "d.pptx"
    prs.save(p)
    return p


def _text(sl, x, y, w, h, text, size=18, color=INK, name="Body", bold=False):
    """Add a word-wrapped text box with one run."""
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tb.name = name
    tb.text_frame.word_wrap = True
    r = tb.text_frame.paragraphs[0].add_run()
    r.text = text
    r.font.size, r.font.bold = Pt(size), bold
    r.font.color.rgb = color
    return tb


def _checks(tmp_path, build):
    """The set of check names QA raises for a one-slide deck."""
    res = qa_render.check_pptx(_deck(tmp_path, build))
    return {f["check"] for fs in res["slides"].values() for f in fs}, res


def _clean(sl):
    """A slide with nothing wrong."""
    _text(sl, 0.8, 0.8, 11, 1.0, "A clear headline that says the point.", size=32, name="Title", bold=True)
    _text(sl, 0.8, 2.2, 11, 2.0, "Each point is short enough to read from the back of the room.", size=18)


def test_clean_slide_has_no_findings(tmp_path):
    checks, res = _checks(tmp_path, _clean)
    assert checks == set() and res["summary"]["total"] == 0 and res["slides"] == {}


def test_off_slide_shape(tmp_path):
    checks, _ = _checks(tmp_path, lambda sl: (_clean(sl), _text(sl, 12.5, 3, 3, 1, "wide", name="Wide")))
    assert "off_slide" in checks


def test_small_text(tmp_path):
    checks, res = _checks(tmp_path, lambda sl: (_clean(sl), _text(sl, 0.8, 5, 6, 0.5, "tiny print", size=8)))
    assert "min_font" in checks
    assert any(f["level"] == "error" for f in res["slides"]["1"])   # 8 pt is 4 under the body floor


def test_short_title_floor(tmp_path):
    checks, _ = _checks(tmp_path, lambda sl: _text(sl, 0.8, 0.8, 11, 1, "A headline", size=24, name="Title"))
    assert "min_font" in checks


def test_low_contrast(tmp_path):
    checks, res = _checks(tmp_path, lambda sl: (_clean(sl), _text(sl, 0.8, 5, 6, 0.5, "pale grey on white", color=RGBColor(0xDD, 0xDD, 0xDD))))
    assert "contrast" in checks


def test_contrast_uses_the_shape_fill(tmp_path):
    def build(sl):
        _clean(sl)
        card = sl.shapes.add_shape(1, Inches(0.8), Inches(5), Inches(5), Inches(1))
        card.fill.solid()
        card.fill.fore_color.rgb = RGBColor(0x10, 0x20, 0x40)
        _text(sl, 1, 5.2, 4, 0.5, "dark ink on a navy card", color=INK)
    assert "contrast" in _checks(tmp_path, build)[0]


def test_leftover_agency_placeholder(tmp_path):
    checks, _ = _checks(tmp_path, lambda sl: (_clean(sl), _text(sl, 0.8, 5, 8, 0.5, "A response from [Agency]")))
    assert "placeholder" in checks


def test_validator_flagged_agency_is_skipped(tmp_path):
    p = _deck(tmp_path, lambda sl: (_clean(sl), _text(sl, 0.8, 5, 8, 0.5, "A response from [Agency]")))
    res = qa_render.check_pptx(p, validation=[{"check": "internal_leak", "slide_no": 1}])
    assert "placeholder" not in {f["check"] for f in res["slides"].get("1", [])}


@pytest.mark.parametrize("word", ["lorem ipsum dolor", "TBD", "XXX"])
def test_other_filler(tmp_path, word):
    assert "placeholder" in _checks(tmp_path, lambda sl: (_clean(sl), _text(sl, 0.8, 5, 8, 0.5, f"Budget: {word}")))[0]


def test_overflowing_frame(tmp_path):
    long = "This sentence keeps going and going so that it cannot fit in a very small frame. " * 6
    checks, _ = _checks(tmp_path, lambda sl: (_clean(sl), _text(sl, 0.8, 4.5, 3, 0.5, long, size=18)))
    assert "text_overflow" in checks


def test_overlapping_text(tmp_path):
    checks, _ = _checks(tmp_path, lambda sl: (_text(sl, 1, 1, 6, 1, "First block of text here", size=20), _text(sl, 2, 1.2, 6, 1, "Second block on top", size=20)))
    assert "text_overlap" in checks


def test_margin(tmp_path):
    assert "margin" in _checks(tmp_path, lambda sl: (_clean(sl), _text(sl, 0.1, 5, 4, 0.5, "hugging the edge")))[0]


def test_hidden_slide_is_not_checked_for_filler(tmp_path):
    p = _deck(tmp_path, lambda sl: (_clean(sl), _text(sl, 0.8, 5, 8, 0.5, "TBD")))
    prs = Presentation(p)
    prs.slides[0]._element.set("show", "0")
    prs.save(p)
    assert "placeholder" not in {f["check"] for fs in qa_render.check_pptx(p)["slides"].values() for f in fs}


def test_pixels_flag_empty_and_corner(tmp_path):
    w, h = 90, 60
    blank = b"\xff\xff\xff" * (w * h)
    assert {f["check"] for f in qa_render._check_pixels(qa_render.ink_shares(w, h, blank))} == {"mostly_empty"}
    px = bytearray(blank)
    for y in range(2, 18):                # a dark block in the top-left corner only
        for x in range(2, 28):
            i = (y * w + x) * 3
            px[i:i + 3] = b"\x10\x10\x10"
    m = qa_render.ink_shares(w, h, bytes(px))
    assert m["h3"][0] > 0 and m["h3"][1] == 0 and "one_corner" in {f["check"] for f in qa_render._check_pixels(m)}
    board = bytearray(blank)              # a checkerboard: content in every third
    for y in range(h):
        for x in range(w):
            if (x // 3 + y // 3) % 2:
                i = (y * w + x) * 3
                board[i:i + 3] = b"\x10\x10\x10"
    full = bytes(board)
    assert qa_render._check_pixels(qa_render.ink_shares(w, h, full)) == []


def test_read_ppm_parses_header_with_comment(tmp_path):
    f = tmp_path / "a.ppm"
    f.write_bytes(b"P6\n# made by a test\n2 1\n255\n" + bytes([1, 2, 3, 4, 5, 6]))
    assert qa_render.read_ppm(f) == (2, 1, bytes([1, 2, 3, 4, 5, 6]))
    f.write_bytes(b"P3\n1 1\n255\n0 0 0")
    with pytest.raises(ValueError):
        qa_render.read_ppm(f)


def test_corrupt_pptx_is_one_error_and_never_raises(tmp_path):
    bad = tmp_path / "bad.pptx"
    bad.write_bytes(b"this is not a zip")
    res = qa_render.run(bad, None, tmp_path / "o")
    assert res["summary"]["errors"] == 1 and res["summary"]["total"] == 1 and "deck" in res["slides"]
    assert json.loads((tmp_path / "o" / "qa.json").read_text())["summary"]["errors"] == 1
    assert qa_render.run(tmp_path / "missing.pptx", tmp_path / "missing.pdf", tmp_path / "o2")["summary"]["errors"] == 1


def test_run_writes_qa_json(tmp_path):
    res = qa_render.run(_deck(tmp_path, lambda sl: (_clean(sl), _text(sl, 0.8, 5, 6, 0.5, "tiny", size=8))), None, tmp_path / "out")
    qa = json.loads((tmp_path / "out" / "qa.json").read_text())
    assert qa["slides"]["1"][0].keys() == {"check", "level", "message"} and qa["summary"]["total"] == res["summary"]["total"] > 0
    assert qa["summary"]["line"].startswith("QA: ")


# ---------------------------------------------------------------- the run_deck hook

def _main(monkeypatch, tmp_path, extra):
    """run_deck.main on the brief fixture with the renderers and LibreOffice stubbed, as the intake tests do."""
    for mod in (run_deck.render_pptx, run_deck.render_web):
        monkeypatch.setattr(mod, "render", lambda *a, **k: tmp_path / "deck.pptx")
    monkeypatch.setattr(run_deck.subprocess, "run", lambda *a, **k: None)
    monkeypatch.setattr(run_deck, "_soffice", lambda: "soffice")
    return run_deck.main(["--brief", str(BRIEF / "brief_object.json"), "--stages", str(BRIEF), "--client", "Lumen Water",
                          "--project", "National awareness campaign", "--out", str(tmp_path / "out"), *extra])


def test_run_deck_runs_qa_and_lists_the_stage(monkeypatch, tmp_path):
    rep = _main(monkeypatch, tmp_path, [])   # the stubbed pptx does not exist: QA reports it and the run still finishes
    assert rep["qa"].startswith("QA: ") and "qa_render" in rep["stages"] and "secs" in rep["stages"]["qa_render"]
    assert (tmp_path / "out" / "qa.json").exists()
    assert json.loads((tmp_path / "out" / "report.json").read_text())["qa"] == rep["qa"]


def test_run_deck_no_qa_skips_it(monkeypatch, tmp_path):
    rep = _main(monkeypatch, tmp_path, ["--no-qa"])
    assert "qa" not in rep and "qa_render" not in rep["stages"] and not (tmp_path / "out" / "qa.json").exists()


@pytest.mark.parametrize("raw", [b"P6 3", b"P6", b"", b"P6 3 2 255", b"P6 # cut off comment", b"P6 x y z "])
def test_read_ppm_truncated_header_raises_not_hangs(tmp_path, raw):
    """A cut-off or garbage PPM header raises ValueError at once (it used to loop forever)."""
    f = tmp_path / "t.ppm"
    f.write_bytes(raw)
    with pytest.raises(ValueError):
        qa_render.read_ppm(f)


def test_draft_prefixed_body_still_flagged(tmp_path):
    """Only the fixed watermark is exempt: a body starting with DRAFT that holds TBD is flagged."""
    checks, _ = _checks(tmp_path, lambda sl: (_clean(sl), _text(sl, 0.8, 5, 8, 0.5, "DRAFT: revenue TBD XXX")))
    assert "placeholder" in checks


def test_real_draft_watermark_is_exempt(tmp_path):
    """The renderer's own DRAFT label raises nothing."""
    checks, _ = _checks(tmp_path, lambda sl: (_clean(sl), _text(sl, 8.7, 0.6, 4, 0.3, "DRAFT · needs a person", size=12, name="Watermark")))
    assert "placeholder" not in checks


def test_table_cell_filler_flagged(tmp_path):
    """Filler inside a table cell on a presented slide is a placeholder error."""
    def build(sl):
        _clean(sl)
        t = sl.shapes.add_table(2, 2, Inches(0.8), Inches(4.5), Inches(6), Inches(1)).table
        for i, txt in enumerate(["Name", "TBD", "[Agency] lorem", "Cost"]):
            c = t.cell(i // 2, i % 2)
            c.text = txt
            for r in c.text_frame.paragraphs[0].runs:
                r.font.size = Pt(14)
    assert "placeholder" in _checks(tmp_path, build)[0]


def test_notes_filler_flagged(tmp_path):
    """Filler in the speaker notes of a presented slide is flagged."""
    def build(sl):
        _clean(sl)
        sl.notes_slide.notes_text_frame.text = "TBD"
    assert "placeholder" in _checks(tmp_path, build)[0]


def test_grouped_shape_filler_flagged(tmp_path):
    """Filler inside a group shape is flagged."""
    def build(sl):
        _clean(sl)
        g = sl.shapes.add_group_shape()
        tb = g.shapes.add_textbox(Inches(0.8), Inches(5), Inches(6), Inches(0.5))
        tb.text_frame.text = "Budget XXX"
        tb.text_frame.paragraphs[0].runs[0].font.size = Pt(18)
    assert "placeholder" in _checks(tmp_path, build)[0]


# ---- the floors by role (title 28, body 12, caption 10, footer 9), read from the theme ----

def test_floors_come_from_the_theme_scale():
    from deck.theme import Scale
    S = Scale()
    assert qa_render.FLOORS == {"title": S.title_min, "body": S.body_floor, "caption": S.caption_floor, "foot": S.foot_floor}


def test_a_caption_frame_may_be_ten_points_but_not_nine(tmp_path):
    ok, _ = _checks(tmp_path, lambda sl: (_clean(sl), _text(sl, 0.8, 5, 4, 0.4, "A LABEL", size=10, name="Caption")))
    assert "min_font" not in ok
    bad, res = _checks(tmp_path, lambda sl: (_clean(sl), _text(sl, 0.8, 5, 4, 0.4, "A LABEL", size=9, name="Caption")))
    assert "min_font" in bad


def test_a_body_frame_at_eleven_points_is_still_flagged(tmp_path):
    checks, _ = _checks(tmp_path, lambda sl: (_clean(sl), _text(sl, 0.8, 5, 6, 0.5, "body copy set too small", size=11)))
    assert "min_font" in checks


def test_footer_and_source_lines_are_exempt_down_to_nine_points(tmp_path):
    """Named allowance: footer furniture may be 9 pt, sits in the margin and is never a body or caption finding."""
    def build(sl):
        _clean(sl)
        _text(sl, 0.65, 7.0, 6, 0.3, "THE ASK", size=9, name="Section")
        _text(sl, 6.2, 7.0, 5.6, 0.3, "Sources [1] [2]", size=9, name="Footnotes")
        _text(sl, 4.5, 0.12, 4.1, 0.3, "CHECK FAILED: word_budget", size=9, color=RGBColor(0xB0, 0x32, 0x25), name="CheckFailed")
    checks, _ = _checks(tmp_path, build)
    assert not checks & {"min_font", "margin"}


def test_footer_under_its_own_floor_is_flagged(tmp_path):
    checks, _ = _checks(tmp_path, lambda sl: (_clean(sl), _text(sl, 0.65, 7.0, 6, 0.3, "THE ASK", size=8, name="Section")))
    assert "min_font" in checks


def test_a_title_at_twenty_eight_points_passes(tmp_path):
    checks, _ = _checks(tmp_path, lambda sl: _text(sl, 0.8, 0.8, 11, 1, "A headline", size=28, name="Title"))
    assert "min_font" not in checks
